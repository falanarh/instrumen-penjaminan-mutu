import csv
from io import StringIO
from pathlib import Path
from urllib.request import urlopen

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill, Border, Side
from openpyxl.utils import get_column_letter


SOURCE_SPREADSHEET_ID = "1RfS8ld3uUiC_5zmZNKqVgOV7zcCrVxd57L2a0AsNUJc"
OUTPUT_DIR = Path("outputs")
OUTPUT_PATH = OUTPUT_DIR / "Instrumen Penjaminan Mutu Pelatihan.xlsx"


def clean(value):
    return (value or "").strip()


def load_source():
    url = f"https://docs.google.com/spreadsheets/d/{SOURCE_SPREADSHEET_ID}/export?format=csv"
    with urlopen(url, timeout=30) as response:
        csv_text = response.read().decode("utf-8-sig")

    rows = list(csv.reader(StringIO(csv_text)))
    metadata = {
        "Nama Pelatihan": clean(rows[3][2]),
        "Angkatan": clean(rows[4][2]),
        "Metode": clean(rows[5][2]),
        "Tanggal Pelatihan": clean(rows[6][2]),
        "Tempat Penyelenggaraan": clean(rows[7][2]),
        "Unit Penyelenggara": clean(rows[8][2]),
        "Tanggal Kegiatan Penjaminan Mutu": clean(rows[9][2]),
    }

    indicators = []
    current_dimension = ""
    seq = 1
    for row in rows[12:]:
        indicator = clean(row[2] if len(row) > 2 else "")
        method = clean(row[3] if len(row) > 3 else "")
        if not indicator or method == "Jumlah":
            continue

        dimension = clean(row[1] if len(row) > 1 else "")
        if dimension:
            current_dimension = dimension

        indicators.append(
            {
                "kode": f"I{seq:03d}",
                "no": clean(row[0] if len(row) > 0 else "") or str(seq),
                "dimensi": current_dimension,
                "indikator": indicator,
                "metode": method,
                "skala": clean(row[4] if len(row) > 4 else "") or "1. ada / 0. tidak",
            }
        )
        seq += 1

    return metadata, indicators


def style_title(ws, cell_range):
    for row in ws[cell_range]:
        for cell in row:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E79")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def style_table(ws, header_row, max_col):
    style_title(ws, f"A{header_row}:{get_column_letter(max_col)}{header_row}")
    thin = Side(style="thin", color="D9E2F3")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for row in ws.iter_rows(min_row=header_row, max_row=ws.max_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def write_metadata(ws, metadata):
    ws["A1"] = "Identitas Pelatihan"
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    row = 2
    for key, value in metadata.items():
        ws.cell(row=row, column=1, value=key).font = Font(bold=True)
        ws.cell(row=row, column=2, value=value)
        row += 1


def write_master(ws, indicators):
    headers = ["Kode Indikator", "No", "Dimensi", "Indikator", "Skala"]
    ws.append(headers)
    for item in indicators:
        ws.append([item["kode"], item["no"], item["dimensi"], item["indikator"], item["skala"]])
    style_table(ws, 1, len(headers))
    ws.freeze_panes = "A2"
    set_widths(ws, [16, 8, 28, 60, 22])


def write_mandiri(ws, metadata, indicators):
    write_metadata(ws, metadata)
    headers = [
        "Kode Indikator",
        "No",
        "Dimensi",
        "Indikator",
        "Skala",
        "Nilai Mandiri",
        "Bukti Dukung",
        "Keterangan",
    ]
    header_row = 10
    for col, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=header)
    for idx, item in enumerate(indicators, start=header_row + 1):
        ws.cell(idx, 1, item["kode"])
        ws.cell(idx, 2, item["no"])
        ws.cell(idx, 3, item["dimensi"])
        ws.cell(idx, 4, item["indikator"])
        ws.cell(idx, 5, item["skala"])
    style_table(ws, header_row, len(headers))
    ws.freeze_panes = "A11"
    set_widths(ws, [16, 8, 28, 60, 22, 16, 42, 42])


def write_verifikasi(ws, metadata, indicators):
    write_metadata(ws, metadata)
    headers = [
        "Kode Indikator",
        "No",
        "Dimensi",
        "Indikator",
        "Skala",
        "Nilai Tim Penilai",
        "Catatan Tim Penilai",
    ]
    header_row = 10
    for col, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=header)
    for idx, item in enumerate(indicators, start=header_row + 1):
        ws.cell(idx, 1, item["kode"])
        ws.cell(idx, 2, item["no"])
        ws.cell(idx, 3, item["dimensi"])
        ws.cell(idx, 4, item["indikator"])
        ws.cell(idx, 5, item["skala"])
    style_table(ws, header_row, len(headers))
    ws.freeze_panes = "A11"
    set_widths(ws, [16, 8, 28, 60, 22, 18, 46])


def write_rekap(ws, metadata, indicators):
    write_metadata(ws, metadata)
    headers = [
        "Kode Indikator",
        "No",
        "Dimensi",
        "Indikator",
        "Skala",
        "Nilai Mandiri",
        "Bukti Dukung",
        "Keterangan Penyelenggara",
        "Nilai Tim Penilai",
        "Catatan Tim Penilai",
        "Nilai Akhir",
    ]
    header_row = 10
    for col, header in enumerate(headers, start=1):
        ws.cell(row=header_row, column=col, value=header)

    for idx, item in enumerate(indicators, start=header_row + 1):
        ws.cell(idx, 1, item["kode"])
        ws.cell(idx, 2, item["no"])
        ws.cell(idx, 3, item["dimensi"])
        ws.cell(idx, 4, item["indikator"])
        ws.cell(idx, 5, item["skala"])
        ws.cell(idx, 6, f"='Penilaian Mandiri Penyelenggara'!F{idx}")
        ws.cell(idx, 7, f"='Penilaian Mandiri Penyelenggara'!G{idx}")
        ws.cell(idx, 8, f"='Penilaian Mandiri Penyelenggara'!H{idx}")
        ws.cell(idx, 9, "")
        ws.cell(idx, 10, "")
        ws.cell(idx, 11, f'=IF(I{idx}<>"",I{idx},F{idx})')

    total_row = header_row + len(indicators) + 2
    ws.cell(total_row, 1, "Jumlah Nilai Akhir")
    ws.cell(total_row, 2, f"=SUM(K{header_row + 1}:K{header_row + len(indicators)})")
    ws.cell(total_row + 1, 1, "Jumlah Indikator Dinilai")
    ws.cell(total_row + 1, 2, f"=COUNT(K{header_row + 1}:K{header_row + len(indicators)})")
    ws.cell(total_row + 2, 1, "Nilai Mutu")
    ws.cell(total_row + 2, 2, f'=IF(B{total_row + 1}=0,"",B{total_row}/B{total_row + 1}*100)')
    for row in range(total_row, total_row + 3):
        ws.cell(row, 1).font = Font(bold=True)

    style_table(ws, header_row, len(headers))
    ws.freeze_panes = "A11"
    set_widths(ws, [16, 8, 28, 60, 22, 16, 42, 42, 18, 46, 16])


def write_panduan(ws, indicator_count):
    rows = [
        ["Instrumen Penjaminan Mutu Pelatihan"],
        [""],
        ["Role", "Penyelenggara mengisi penilaian mandiri melalui Google Form. Tim Penilai mengisi penilaian lanjutan langsung di Sheet."],
        ["Alur 1", "Isi tab Penilaian Mandiri Penyelenggara."],
        ["Alur 2", "Isi Nilai Mandiri, Bukti Dukung, dan Keterangan bila diperlukan."],
        ["Alur 3", "Tim Penilai melihat Rekap, lalu mengisi kolom Nilai Tim Penilai dan Catatan Tim Penilai langsung di tab Rekap."],
        ["Alur 4", "Tab Rekap otomatis mengambil nilai mandiri; nilai akhir memakai nilai tim penilai jika tersedia."],
        [""],
        ["Aturan Nilai", "Gunakan angka 0 sampai 1. Gunakan tanda - untuk indikator tidak berlaku dan jelaskan di Keterangan/Catatan."],
        ["Bukti Dukung", "Dapat berupa link gojags, kalender, evaluasi, dokumen, file soft copy, atau ringkasan hasil pengamatan lapangan."],
        ["Catatan", "Workbook ini mendukung satu Google Form untuk Penilaian Mandiri. Penilaian Tim Penilai dilakukan langsung di tab Rekap."],
        ["Jumlah Indikator", indicator_count],
    ]
    for row in rows:
        ws.append(row)
    ws["A1"].font = Font(bold=True, size=14, color="1F4E79")
    set_widths(ws, [24, 110])
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def set_widths(ws, widths):
    for idx, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width


def main():
    metadata, indicators = load_source()
    wb = Workbook()
    wb.remove(wb.active)

    write_panduan(wb.create_sheet("Panduan"), len(indicators))
    write_master(wb.create_sheet("Master Indikator"), indicators)
    write_mandiri(wb.create_sheet("Penilaian Mandiri Penyelenggara"), metadata, indicators)
    write_rekap(wb.create_sheet("Rekap"), metadata, indicators)

    OUTPUT_DIR.mkdir(exist_ok=True)
    wb.save(OUTPUT_PATH)
    print(OUTPUT_PATH.resolve())


if __name__ == "__main__":
    main()
