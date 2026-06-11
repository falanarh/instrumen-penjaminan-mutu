from pathlib import Path
import shutil
import time

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "outputs" / "sosialisasi_instrumen_penjaminan_mutu.pptx"
BACKUP = ROOT / "outputs" / f"sosialisasi_instrumen_penjaminan_mutu_before_links_slide_{time.strftime('%Y%m%d%H%M%S')}.pptx"
FALLBACK = ROOT / "outputs" / f"sosialisasi_instrumen_penjaminan_mutu_with_links_{time.strftime('%Y%m%d%H%M%S')}.pptx"

INK = RGBColor(23, 32, 51)
MUTED = RGBColor(102, 112, 133)
BLUE = RGBColor(31, 78, 121)
BLUE_BG = RGBColor(220, 235, 250)
GREEN = RGBColor(46, 125, 50)
GREEN_BG = RGBColor(217, 234, 211)
ORANGE = RGBColor(198, 89, 17)
ORANGE_BG = RGBColor(252, 228, 214)
WHITE = RGBColor(255, 255, 255)
GRAY = RGBColor(152, 162, 179)

FORM_URL = "https://docs.google.com/forms/d/1uHzZrUlV2MuCAxA6DwmzI0q-IoAZJhlY5FUHOKEWH44/edit"
SHEET_URL = "https://docs.google.com/spreadsheets/d/1nJGiUrdr-QCb9GoltjmxMS9BK_i1eGwwx7VjDTLJ8Z8/edit"


def add_text(slide, text, x, y, w, h, size, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_link_card(slide, x, y, title, url, color, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(5.45), Inches(1.45))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = color
    shape.line.width = Pt(1.2)
    add_text(slide, title, x + 0.28, y + 0.24, 4.8, 0.3, 18, color, True)
    box = add_text(slide, url, x + 0.28, y + 0.72, 4.85, 0.34, 9.8, BLUE)
    box.click_action.hyperlink.address = url


def add_change(slide, n, text, y):
    circ = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.95), Inches(y), Inches(0.33), Inches(0.33))
    circ.fill.solid()
    circ.fill.fore_color.rgb = ORANGE
    circ.line.fill.background()
    add_text(slide, str(n), 0.95, y + 0.075, 0.33, 0.11, 8.5, WHITE, True, PP_ALIGN.CENTER)
    add_text(slide, text, 1.45, y + 0.025, 10.6, 0.3, 13.5, INK)


def renumber_footers(prs):
    for idx, slide in enumerate(prs.slides, start=1):
        expected = f"{idx:02d}"
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text.strip()
            if len(text) == 2 and text.isdigit():
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT
                run = p.add_run()
                run.text = expected
                run.font.name = "Aptos"
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = GRAY


def main():
    if not PPTX.exists():
        raise FileNotFoundError(PPTX)

    shutil.copy2(PPTX, BACKUP)
    prs = Presentation(str(PPTX))
    slide = prs.slides.add_slide(prs.slide_layouts[0])

    add_text(slide, "Link operasional", 0.7, 0.55, 8.8, 0.48, 28, INK, True)
    add_text(slide, "Gunakan dua tautan berikut sesuai peran masing-masing.", 0.72, 1.15, 9.2, 0.32, 12.5, MUTED)
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.72), Inches(1.1), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = BLUE
    line.line.fill.background()

    add_link_card(slide, 0.95, 2.15, "Form Penilaian Mandiri", FORM_URL, BLUE, BLUE_BG)
    add_link_card(slide, 6.85, 2.15, "Google Sheet Tim Penilai", SHEET_URL, GREEN, GREEN_BG)

    add_text(slide, "Ringkasan perubahan terakhir", 0.95, 4.25, 5.8, 0.35, 18, INK, True)
    add_change(slide, 1, "Isian Angkatan dihapus; penilaian mutu dilakukan satu kali per penyelenggaraan pelatihan.", 4.85)
    add_change(slide, 2, "Metode pemeriksaan dihapus dari Form dan Sheet.", 5.38)
    add_change(slide, 3, "Bukti dukung dapat berupa upload file atau tautan dokumen publik.", 5.91)
    add_change(slide, 4, "Skala penilaian mandiri: 0, 1, atau -.", 6.44)

    add_text(slide, "Instrumen Penjaminan Mutu Pelatihan", 0.7, 7.03, 4.2, 0.22, 8.5, GRAY)
    add_text(slide, f"{len(prs.slides):02d}", 12.15, 7.0, 0.45, 0.25, 9, GRAY, True, PP_ALIGN.RIGHT)

    renumber_footers(prs)
    try:
        prs.save(str(PPTX))
        output = PPTX
    except PermissionError:
        prs.save(str(FALLBACK))
        output = FALLBACK

    print(f"updated={output}")
    print(f"backup={BACKUP}")


if __name__ == "__main__":
    main()
