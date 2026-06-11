from copy import deepcopy
from pathlib import Path
import shutil
import time

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[1]
PPTX = ROOT / "outputs" / "sosialisasi_instrumen_penjaminan_mutu.pptx"
BACKUP = ROOT / "outputs" / f"sosialisasi_instrumen_penjaminan_mutu_before_platform_slide_{time.strftime('%Y%m%d%H%M%S')}.pptx"

INK = RGBColor(23, 32, 51)
MUTED = RGBColor(102, 112, 133)
BLUE = RGBColor(31, 78, 121)
BLUE_BG = RGBColor(220, 235, 250)
GREEN = RGBColor(46, 125, 50)
GREEN_BG = RGBColor(217, 234, 211)
LINE = RGBColor(208, 213, 221)
WHITE = RGBColor(255, 255, 255)


def add_text(slide, text, x, y, w, h, size, color=INK, bold=False, align=None):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.clear()
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    p = tf.paragraphs[0]
    if align:
        p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = "Aptos"
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    return box


def add_line(slide, x, y, w, color=BLUE, width=2):
    line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(0.03))
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.fill.background()


def add_platform_card(slide, x, title, subtitle, color, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.25), Inches(4.55), Inches(1.65))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = color
    shape.line.width = Pt(1.4)
    add_text(slide, title, x + 0.28, 2.55, 3.9, 0.35, 22, color, True)
    add_text(slide, subtitle, x + 0.28, 3.05, 3.85, 0.45, 15, INK)


def add_role_row(slide, y, role, form_text, sheet_text):
    add_text(slide, role, 1.05, y + 0.08, 2.25, 0.28, 16, INK, True)
    for x, txt, color in [(4.05, form_text, BLUE), (8.05, sheet_text, GREEN)]:
        fill = BLUE_BG if color == BLUE else GREEN_BG
        box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(2.65), Inches(0.55))
        box.fill.solid()
        box.fill.fore_color.rgb = fill
        box.line.fill.background()
        add_text(slide, txt, x + 0.12, y + 0.15, 2.42, 0.2, 12.5, color, True, PP_ALIGN.CENTER)


def insert_slide(prs, slide, index):
    sld_id_lst = prs.slides._sldIdLst
    sld_id = sld_id_lst[-1]
    sld_id_lst.remove(sld_id)
    sld_id_lst.insert(index, sld_id)


def renumber_footers(prs):
    for idx, slide in enumerate(prs.slides, start=1):
        expected = f"{idx:02d}"
        for shape in slide.shapes:
            if not getattr(shape, "has_text_frame", False):
                continue
            text = shape.text.strip()
            if len(text) == 2 and text.isdigit():
                shape.text_frame.clear()
                p = shape.text_frame.paragraphs[0]
                p.alignment = PP_ALIGN.RIGHT
                run = p.add_run()
                run.text = expected
                run.font.name = "Aptos"
                run.font.size = Pt(9)
                run.font.bold = True
                run.font.color.rgb = RGBColor(152, 162, 179)


def main():
    if not PPTX.exists():
        raise FileNotFoundError(PPTX)

    shutil.copy2(PPTX, BACKUP)
    prs = Presentation(str(PPTX))
    blank = prs.slide_layouts[0]
    slide = prs.slides.add_slide(blank)

    add_text(slide, "Dua platform yang digunakan", 0.7, 0.55, 8.8, 0.48, 28, INK, True)
    add_text(
        slide,
        "Peran dibuat sederhana: penyelenggara memakai Form dan Sheet, sedangkan tim penilai bekerja hanya di Google Sheets.",
        0.72,
        1.15,
        9.2,
        0.36,
        12.5,
        MUTED,
    )
    add_line(slide, 0.7, 1.72, 1.1)

    add_platform_card(slide, 1.05, "Google Form", "Tempat penyelenggara mengirim atau mengedit jawaban.", BLUE, BLUE_BG)
    add_platform_card(slide, 6.95, "Google Sheets", "Tempat monitoring progress dan penilaian indikator.", GREEN, GREEN_BG)

    add_text(slide, "Siapa menggunakan apa?", 1.05, 4.45, 4.8, 0.35, 20, INK, True)
    add_role_row(slide, 5.05, "Penyelenggara", "Mengisi Form", "Memantau Progress")
    add_role_row(slide, 5.88, "Tim Penilai", "-", "Menilai di Sheet")

    add_text(
        slide,
        "Intinya: penyelenggara punya jalur input dan monitoring; penilai cukup fokus di satu platform kerja.",
        1.05,
        6.65,
        9.7,
        0.32,
        15,
        BLUE,
        True,
    )
    add_text(slide, "Instrumen Penjaminan Mutu Pelatihan", 0.7, 7.03, 4.2, 0.22, 8.5, RGBColor(152, 162, 179))
    add_text(slide, "02", 12.15, 7.0, 0.45, 0.25, 9, RGBColor(152, 162, 179), True, PP_ALIGN.RIGHT)

    insert_slide(prs, slide, 1)
    renumber_footers(prs)
    prs.save(str(PPTX))
    print(f"updated={PPTX}")
    print(f"backup={BACKUP}")


if __name__ == "__main__":
    main()
